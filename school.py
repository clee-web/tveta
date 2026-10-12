"""Helpers for classes, teachers, exams, attendance and QR codes."""

import base64
import io
import json
import re

import qrcode
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


# ------------------------------------------------------------------
# Teachers
# ------------------------------------------------------------------
def generate_teacher_username(db, full_name):
    """firstname.lastname, with a number appended if it is taken."""
    parts = re.findall(r"[A-Za-z0-9]+", (full_name or "").lower())
    base = ".".join(parts[:2]) if parts else "teacher"
    base = re.sub(r"[^a-z0-9.]", "", base) or "teacher"

    username = base
    n = 1
    while db.execute(
        "SELECT 1 FROM teachers WHERE username = ?", (username,)
    ).fetchone():
        n += 1
        username = f"{base}{n}"
    return username


# ------------------------------------------------------------------
# Classes / students
# ------------------------------------------------------------------
def students_in_class(db, cls):
    """Admitted students matching a class (course + year + study mode)."""
    return db.execute(
        """SELECT * FROM applications
           WHERE admission_number IS NOT NULL
             AND course = ? AND academic_year = ? AND study_mode = ?
           ORDER BY full_name COLLATE NOCASE""",
        (cls["course"], cls["academic_year"], cls["study_mode"]),
    ).fetchall()


def class_for_student(db, student):
    return db.execute(
        """SELECT * FROM classes
           WHERE course = ? AND academic_year = ? AND study_mode = ?""",
        (student["course"], student["academic_year"], student["study_mode"]),
    ).fetchone()


def is_student_in_class(student, cls):
    return (
        student["course"] == cls["course"]
        and student["academic_year"] == cls["academic_year"]
        and student["study_mode"] == cls["study_mode"]
    )


# ------------------------------------------------------------------
# Grading
# ------------------------------------------------------------------
def compute_grade(score, max_score=100):
    if score is None:
        return None
    try:
        pct = (float(score) / float(max_score)) * 100 if max_score else 0
    except (TypeError, ValueError, ZeroDivisionError):
        return None
    if pct >= 80:
        return "A"
    if pct >= 70:
        return "B"
    if pct >= 60:
        return "C"
    if pct >= 50:
        return "D"
    if pct >= 40:
        return "E"
    return "F"


# ------------------------------------------------------------------
# QR codes
# ------------------------------------------------------------------
def qr_png_bytes(text):
    """Return an in-memory PNG of a QR code for the given text."""
    img = qrcode.make(text or "")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def qr_data_uri(text):
    return "data:image/png;base64," + base64.b64encode(qr_png_bytes(text).getvalue()).decode("ascii")


def student_qr_payload(student):
    """QR payload carrying identifying details so a scan reveals who the student is."""
    return json.dumps(
        {
            "type": "amuaa_student",
            "admission_number": student["admission_number"] or student["id"],
            "full_name": student["full_name"],
            "course": student["course"],
            "academic_year": student["academic_year"] or "",
            "study_mode": student["study_mode"] or "",
        },
        separators=(",", ":"),
        ensure_ascii=False,
    )


def admission_from_qr(raw):
    """Return an admission number from a JSON QR payload or from plain text."""
    text = (raw or "").strip()
    if text.startswith("{"):
        try:
            data = json.loads(text)
            return str(data.get("admission_number") or data.get("adm") or text).strip()
        except (ValueError, TypeError):
            return text
    return text


# ------------------------------------------------------------------
# Excel exam templates
# ------------------------------------------------------------------
_HEADER_FILL = PatternFill("solid", fgColor="762B3A")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def build_exam_workbook(exam, students):
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    ws.append(["Admission No", "Student Name", f"Score (out of {float(exam['max_score']):.0f})"])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for s in students:
        ws.append([s["admission_number"], s["full_name"], None])

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 20
    ws.freeze_panes = "A2"
    return wb


def read_exam_workbook(file_stream):
    """Return a list of (admission_number, score) tuples from an uploaded sheet."""
    wb = load_workbook(file_stream, data_only=True)
    ws = wb.active

    rows = []
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        if index == 0:  # header
            continue
        if not row:
            continue
        adm = row[0] if len(row) > 0 else None
        score = row[2] if len(row) > 2 else None
        if adm is None or str(adm).strip() == "":
            continue
        rows.append((str(adm).strip(), score))
    return rows
