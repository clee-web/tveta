import csv
import io
from datetime import datetime

from flask import Response, request, redirect, url_for, flash

from database import get_db
from utils import login_required
from ._bp import admin_bp


@admin_bp.route("/export/csv")
@login_required
def export_csv():
    rows = _fetch_applications()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(_headers())
    for r in rows:
        writer.writerow(_row(r))

    filename = f"amuaa_applications_{_filter()}_{datetime.utcnow():%Y%m%d}.csv"
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@admin_bp.route("/export/xlsx")
@login_required
def export_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        flash("Install openpyxl first:  pip install openpyxl", "danger")
        return redirect(url_for("admin.applications"))

    rows = _fetch_applications()

    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append(_headers())

    header_fill = PatternFill("solid", fgColor="762B3A")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, _ in enumerate(_headers(), start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append(_row(r))

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max_len, 42)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"amuaa_applications_{_filter()}_{datetime.utcnow():%Y%m%d}.xlsx"
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ------------------------------------------------------------------
# Internal helpers
# ------------------------------------------------------------------
def _filter():
    return request.args.get("status", "all")


def _headers():
    return [
        "Reference", "Surname", "First name", "Middle name", "Full name",
        "Email", "Mobile", "Nationality", "Date of birth", "Medical condition",
        "Course", "Study mode", "Start date", "End date",
        "Next of kin", "Next of kin relationship",
        "Sponsor", "Sponsor organisation", "Sponsor contact person",
        "Previous education", "Status", "Notes", "Submitted",
    ]


def _row(r):
    return [
        r["id"][:8].upper(),
        r["surname"] or "", r["first_name"] or "", r["middle_name"] or "",
        r["full_name"], r["email"], r["phone"],
        r["nationality"] or "", r["date_of_birth"] or "",
        r["medical_condition"] or "",
        r["course"], r["study_mode"],
        r["start_date"] or "", r["end_date"] or "",
        r["next_of_kin_name"] or "", r["next_of_kin_relationship"] or "",
        r["sponsor_name"] or "", r["sponsor_organisation"] or "",
        r["sponsor_contact_person"] or "",
        r["previous_education"] or "",
        r["status"], r["notes"] or "",
        r["created_at"][:19].replace("T", " "),
    ]


def _fetch_applications():
    status_filter = _filter()
    search = request.args.get("q", "").strip()
    query = "SELECT * FROM applications"
    params = []
    clauses = []
    if status_filter in ("pending", "verified", "rejected"):
        clauses.append("status = ?")
        params.append(status_filter)
    if search:
        clauses.append("(full_name LIKE ? OR email LIKE ? OR id LIKE ? OR course LIKE ?)")
        like = f"%{search}%"
        params.extend([like, like, like, like])
    if clauses:
        query += " WHERE " + " AND ".join(clauses)
    query += " ORDER BY created_at DESC"

    with get_db() as db:
        return db.execute(query, params).fetchall()